"""Bounded spreadsheet preview; never evaluates formulas or saves uploaded files."""
import base64
import binascii
import csv
import html
import io
import re
from itertools import chain, islice
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException


def ignored_header(value):
    normalized = re.sub(r'\s+', ' ', value.lower().replace('_', ' ').replace('-', ' ')).strip()
    return normalized in ('device id', 'o.o', 'o.o value')


def pasted_markdown_rows(text):
    """Read a copied Markdown table without treating its separator as an item."""
    def cells(line):
        parts=re.split(r'(?<!\\)\|',line.strip().strip('|'))
        return [html.unescape(re.sub(r'\\([|*`_])',r'\1',part)).strip() for part in parts]
    lines=text.splitlines()
    for index in range(1,min(2,len(lines))):
        separator=cells(lines[index])
        if len(separator)>1 and all(re.fullmatch(r':?-{3,}:?',part) for part in separator) and '|' in lines[index-1]:
            first=cells(lines[index-1])
            data=[]
            for line in lines[index+1:]:
                if not line.strip():continue
                if '|' not in line:break
                data.append(cells(line))
            known={'brand','model','description','serial number','quantity','location','notes','channel','channel name','name','ip','ip address','multicast ip','port','source','venue','device'}
            normalized={header_name(value) for value in first if value.strip()}
            looks_like_data=not normalized.intersection(known) and (
                any(re.fullmatch(r'(?:\d{1,3}\.){3}\d{1,3}',value.strip()) for value in first)
                or sum(bool(re.fullmatch(r'-?\d+',value.strip())) for value in first)>=2
            )
            rows=[first,*data]
            if looks_like_data:
                width=max(map(len,rows))
                rows=[[f'Column {number}' for number in range(1,width+1)],*rows]
                return rows, True
            return rows, False
    return None


def header_name(value):
    return re.sub(r'\s+', ' ', re.sub(r'[_-]+', ' ', str(value or '').lower())).strip()


def detect_header(rows, aliases=()):
    """Score candidates against familiar headings and following data, including custom sheets."""
    known={header_name(v) for v in ('brand','manufacturer','make','model','model number','description','item','serial number','serial','s/n','quantity','qty','quantity in stock','location','notes','venue','device name','channel name','ip','ip address','multicast ip','vlan','port','id','scala id','orientation','size','monitor model','monitor brand','total','total count','deck 6','deck 2 port','deck 2 stbd')}
    known.update(header_name(v) for v in aliases if isinstance(v,str))
    candidates=[]
    for index,row in enumerate(rows[:50]):
        values=[header_name(v) for v in row if str(v or '').strip()]
        if not values:continue
        matched=sum(value in known for value in values)
        placeholders=sum(bool(re.fullmatch(r'column\s*\d+',value)) for value in values)
        text=sum(not re.fullmatch(r'[\d.]+',value) for value in values)
        following=rows[index+1:index+4]
        populated=sum(bool(any(str(v or '').strip() for v in data)) for data in following)
        score=matched*6+min(len(values),12)+text/len(values)+populated-placeholders*5
        if len(values)==1:score-=8
        candidates.append((score,-index,matched,index+1))
    if not candidates:return 1,'low'
    best=max(candidates)
    return best[3], 'high' if best[2]>=2 else 'low'


def excel_cell_text(cell):
    value=cell.value
    # Excel identifiers often use a numeric cell with a zero-padding display format.
    if isinstance(value,(int,float)) and not isinstance(value,bool) and float(value).is_integer():
        number=int(value)
        if re.fullmatch(r'0{2,}',cell.number_format or '') and number>=0:return str(number).zfill(len(cell.number_format))
        return number
    return value


def preview(payload):
    if not isinstance(payload, dict):
        raise ValueError('Provide a spreadsheet file.')
    filename = payload.get('filename', '')
    if not isinstance(filename, str) or Path(filename).suffix.lower() not in ('.xlsx', '.csv'):
        raise ValueError('Use .xlsx or .csv. Save older .xls files as .xlsx in Excel first.')
    try:
        raw = base64.b64decode(payload.get('content', ''), validate=True)
    except (binascii.Error, TypeError, ValueError):
        raise ValueError('Invalid file content.') from None
    if len(raw) > 5_000_000 or (not raw and filename.lower().endswith('.xlsx')):
        raise ValueError('Choose a spreadsheet between 1 byte and 5 MB.')
    pasted=payload.get('pasted') is True
    header_row = 1 if pasted else payload.get('header_row', 1)
    if isinstance(header_row, bool) or not isinstance(header_row, int) or not 1 <= header_row <= 50:
        raise ValueError('Header row must be a number from 1 to 50.')
    book = None
    sheets, sheet = [], ''
    try:
        if filename.lower().endswith('.xlsx'):
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist()) > 500 or sum(f.file_size for f in archive.infolist()) > 30_000_000:
                    raise ValueError('This workbook is too large when expanded. Export a smaller sheet as CSV.')
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
            sheets = book.sheetnames
            if payload.get('sheets_only'):
                return {'sheets': sheets}
            sheet = payload.get('sheet') or sheets[0]
            if sheet not in sheets:
                raise ValueError('Choose a worksheet from this workbook.')
            worksheet = book[sheet]
            worksheet.reset_dimensions()
            source = ([excel_cell_text(cell) for cell in row] for row in worksheet.iter_rows(values_only=False))
        else:
            try:
                text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                raise ValueError('Save your CSV as CSV UTF-8 in Excel, then try again.') from None
            delimiter=payload.get('delimiter')
            if delimiter is not None and delimiter not in (',',';','\t','|'):
                raise ValueError('Choose a supported spreadsheet separator.')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            markdown=pasted_markdown_rows(text) if payload.get('pasted') is True else None
            generated_headers=bool(markdown and markdown[1])
            source = markdown[0] if markdown is not None else csv.reader(io.StringIO(text), dialect, **({'delimiter':delimiter} if delimiter is not None else {}))
        confidence=None
        header_candidates=[]
        if payload.get('detect_header') and not pasted:
            source=iter(source)
            first=list(islice(source,53))
            header_row,confidence=detect_header(first,payload.get('header_aliases',[]) + payload.get('keep_headers',[]))
            header_candidates=[{'number':i+1,'values':[str(v or '') for v in row]} for i,row in enumerate(first[:50]) if any(str(v or '').strip() for v in row)]
            source=chain(first,source)
        headers, rows, row_numbers, characters = [], [], [], 0
        kept_columns, ignored_columns, column_count = [], [], None
        for number, raw_row in enumerate(source, 1):
            if number < header_row:
                continue
            # Restrict raw rows too, so millions of empty rows cannot consume unlimited work.
            if number > 20050:
                raise ValueError('Too many spreadsheet rows. Remove unused rows and try again.')
            values = []
            for value in raw_row:
                if isinstance(value, float) and value.is_integer():
                    value = int(value)
                value = '' if value is None else str(value).strip()
                if len(value) > 2000:
                    raise ValueError(f'Row {number} contains a cell longer than 2,000 characters.')
                values.append(value)
            while values and not values[-1]:
                values.pop()
            if len(values) > 50:
                raise ValueError('Import supports up to 50 columns. Export only the device columns.')
            characters += sum(map(len, values))
            if characters > 2_000_000:
                raise ValueError('Spreadsheet text is too large. Split the file into smaller imports.')
            if pasted and number==1 and not values:
                raise ValueError('The first pasted row must contain column headings. Copy the headings with your Excel cells.')
            if column_count is None:
                if payload.get('auto_header') and not pasted:
                    recognized = {'brand','manufacturer','make','model','model number','description','item','serial number','serial','s/n','quantity','qty','quantity in stock','location','notes'}
                    recognized.update(payload.get('header_aliases', []))
                    count = sum(re.sub(r'[_-]+', ' ', v.lower()).strip() in recognized for v in values)
                    if count < payload.get('minimum_headers',3):
                        if number >= 50:
                            raise ValueError('Could not identify an inventory header within the first 50 rows.')
                        continue
                if not values:
                    continue
                header_row=number
                column_count = len(values)
                keep = {re.sub(r'[_-]+', ' ', h.lower()).strip() for h in payload.get('keep_headers', []) if isinstance(h,str)}
                kept_columns = [i for i, value in enumerate(values) if payload.get('pasted') is True or not ignored_header(value) or re.sub(r'[_-]+',' ',value.lower()).strip() in keep]
                ignored_columns = [value for i,value in enumerate(values) if i not in kept_columns]
                headers = [values[i] or f'Column {i + 1}' for i in kept_columns]
            else:
                if len(values) > column_count:
                    if not pasted:
                        raise ValueError(f'Row {number} has more columns than the chosen header row.')
                    # Excel omits empty cells at the end of copied rows. If a later
                    # pasted row contains data in one of those cells, retain it as a
                    # temporary column instead of blocking the whole import.
                    previous_count=column_count
                    column_count=len(values)
                    kept_columns.extend(range(previous_count,column_count))
                    headers.extend(f'Column {index + 1}' for index in range(previous_count,column_count))
                    for saved_row in rows:
                        saved_row.extend(['']*(column_count-len(saved_row)))
                    generated_headers=True
                values += [''] * (column_count - len(values))
                values = ['' if values[i].lower() == 'o.o' else values[i] for i in kept_columns]
                # Rows containing only IDs or placeholders do not describe a device.
                if not any(values):
                    continue
                if len(rows) >= 10000:
                    raise ValueError('Import supports up to 10,000 devices per file.')
                rows.append(values)
                row_numbers.append(number)
        if payload.get('auto_header') and not headers:
            raise ValueError('Could not identify an inventory header within the first 50 rows.')
        return {'headers': headers, 'rows': rows, 'row_numbers': row_numbers, 'sheets': sheets, 'sheet': sheet, 'ignored_columns': ignored_columns, 'header_row':header_row,'header_confidence':confidence, 'header_candidates':header_candidates, 'generated_headers':locals().get('generated_headers',False),'column_samples':[[row[i] for row in rows[:4]] for i in range(len(headers))]}
    except (BadZipFile, InvalidFileException, ParseError, DefusedXmlException, KeyError, csv.Error, OSError):
        raise ValueError('Could not read this spreadsheet. Save a fresh .xlsx or CSV UTF-8 copy in Excel.') from None
    finally:
        if book is not None:
            book.close()
