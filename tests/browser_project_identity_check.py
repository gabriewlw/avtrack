from browser_helpers import open_workspace, open_import_settings, set_import_confirmation, open_import_matching, start_import_review, choose_new_inventory_import
"""Official AVTrack names reach actual downloads for every inventory and format."""
import io,json,shutil,sys,tempfile,threading
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

with tempfile.TemporaryDirectory() as tmp:
    app.DB_PATH=Path(tmp)/'inventory.sqlite3'
    app.save_device(dict(name='AMX Control',ip='10.1.2.3',venue='Studio'))
    app.save_device(dict(name='Crew TV',ip='239.1.1.1',port=1234,record_type='iptv'))
    inventory=app.save_equipment_inventory(dict(name='Spare TVs',layout=dict(app.DEFAULT_PROFILE,show_status=False)))
    app.save_equipment(dict(inventory_id=inventory['id'],model='OLED55C4',brand='LG',quantity=2,location='Locker'))
    with app.connect() as connection:connection.commit()
    server = app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(executable_path=shutil.which('chromium'))
            page=browser.new_page(viewport={'width':1440,'height':1000});errors=[]
            page.on('pageerror',lambda error:errors.append(str(error)))
            page.goto(f'http://127.0.0.1:{server.server_port}/')
            page.wait_for_function("() => !document.getElementById('refresh').disabled")
            assert page.locator('.brand-logo').inner_text()=='AVTRACK'
            assert page.locator('meta[name=application-name]').get_attribute('content')=='avtrack'
            assert 'broadcasthub' not in page.content().lower()
            for tab,prefix,panel in [('device','av-devices','#inventory'),('iptv','iptv-channels','#inventory'),('equipment','spare-tvs','#equipment-panel')]:
                open_workspace(page, 'equipment-tab' if tab=='equipment' else 'iptv-tab' if tab=='iptv' else 'device-tab')
                if tab=='equipment':
                    page.wait_for_function("() => !document.getElementById('equipment-inventory-select').disabled")
                    page.locator('#equipment-inventory-select').select_option(str(inventory['id']))
                    page.wait_for_function("() => !document.getElementById('equipment-inventory-select').disabled")
                for extension,label in [('json','JSON'),('csv','CSV'),('xlsx','XLSX'),('pdf','PDF')]:
                    menu=page.locator(panel+' .export-dropdown');menu.locator('summary').click()
                    with page.expect_download() as result:
                        menu.get_by_role('link',name=label+' ↓',exact=True).click()
                    download=result.value
                    assert download.suggested_filename==f'avtrack-{prefix}.{extension}',download.suggested_filename
                    raw=Path(download.path()).read_bytes()
                    if extension=='json':
                        records=json.loads(raw)['equipment' if tab=='equipment' else 'devices']
                        assert len(records)==1
                    elif extension=='xlsx':
                        book=load_workbook(io.BytesIO(raw))
                        assert book.properties.creator=='avtrack' and len(book.sheetnames)==1
                    elif extension=='pdf':assert raw.startswith(b'%PDF-') and b'/Author (avtrack)' in raw
                    else:assert ('OLED55C4' if tab=='equipment' else 'Crew TV' if tab=='iptv' else 'AMX Control') in raw.decode('utf-8-sig')
            assert not errors,errors
            browser.close()
    finally:server.shutdown();server.server_close()
print('Official AVTrack identity and all 12 browser downloads with scoped records passed.')
