"""Mobile operations: navigation, version, CRUD, confirmations, scoped exports and deletion.
The former monolith's import/data-rule cases are exercised by the dedicated
import, DHCP, CSV, IPTV, inventory and export suites in the full audit.
"""
import csv, io, json, re, shutil, subprocess, sys, tempfile, threading
from pathlib import Path
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import app
from browser_helpers import open_workspace
from playwright.sync_api import sync_playwright, expect

with tempfile.TemporaryDirectory() as tmp:
    app.DB_PATH=Path(tmp)/'mobile.sqlite3'
    app.save_device(dict(name='Camera A',ip='192.0.2.10',venue='Theater',discipline='Video',vlan=120))
    app.save_device(dict(name='Camera B',ip='192.0.2.11',venue='Theater',discipline='Video',vlan=120))
    app.save_device(dict(name='Audio desk',ip='192.0.2.20',venue='Studio',discipline='Audio'))
    app.save_device(dict(record_type='iptv',name='News stream',ip='239.252.1.1',port=1234,channel_source='Satellite'))
    app.save_device(dict(record_type='iptv',name='Ship stream',ip='239.252.1.2',port=1234,channel_source='Onboard'))
    inv=app.save_equipment_inventory(dict(name='Mobile stock'))
    other=app.save_equipment_inventory(dict(name='Retained inventory'))
    app.save_equipment(dict(inventory_id=other['id'],description='Keep me',quantity=1))
    server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(executable_path=shutil.which('chromium') or None)
            context=browser.new_context(viewport={'width':390,'height':844},is_mobile=True,has_touch=True,device_scale_factor=2)
            page=context.new_page();errors=[]
            page.on('pageerror',lambda error:errors.append(str(error)))
            origin=f'http://127.0.0.1:{server.server_port}'
            page.goto(origin)
            page.wait_for_function("()=>!document.getElementById('refresh').disabled")
            version=page.locator('[aria-label="AVTrack version"]').inner_text()
            count=subprocess.check_output(['git','rev-list','--count','HEAD'],cwd=app.ROOT).decode().strip()
            assert version=='v0.0.'+count and re.fullmatch(r'v0\.0\.\d+',version)
            toggle=page.locator('#nav-toggle')
            assert toggle.is_visible() and toggle.get_attribute('aria-expanded')=='false'
            toggle.focus();toggle.press('Enter')
            assert page.get_by_role('navigation',name='Main navigation').is_visible()
            page.keyboard.press('Escape')
            assert toggle.get_attribute('aria-expanded')=='false'
            assert toggle.evaluate('(n)=>n===document.activeElement')
            open_workspace(page,'device-tab')
            expect(page.locator('.device-row')).to_have_count(3)
            assert toggle.get_attribute('aria-expanded')=='false'
            page.get_by_role('button',name='Confirm IP for Camera A',exact=True).click()
            page.get_by_role('button',name='IP confirmed for Camera A',exact=True).wait_for()
            page.reload()
            page.get_by_role('button',name='IP confirmed for Camera A',exact=True).wait_for()
            page.get_by_role('button',name='Edit Camera A',exact=True).click()
            page.locator('#device-form [name=ip]').fill('192.0.2.999')
            page.locator('#save-device').click() if page.locator('#save-device').count() else page.get_by_role('button',name='Save changes',exact=True).click()
            expect(page.locator('#form-error')).to_contain_text('valid IPv4')
            page.locator('#device-form [name=ip]').fill('192.0.2.10')
            page.locator('#device-form [name=name]').fill('Camera Alpha')
            page.get_by_role('button',name='Save changes',exact=True).click()
            page.locator('#device-dialog').wait_for(state='hidden')
            page.get_by_role('button',name='IP confirmed for Camera Alpha',exact=True).wait_for()

            def download(panel,label,scope):
                menu=page.locator(panel+' .export-dropdown')
                if not menu.evaluate('(n)=>n.open'):menu.locator('summary').click()
                menu.locator('.export-scope').select_option(scope)
                box=menu.locator('.export-links').bounding_box()
                assert box['x']>=0 and box['x']+box['width']<=page.viewport_size['width']+1
                with page.expect_download() as result:
                    menu.get_by_role('link',name=label+' ↓',exact=True).click()
                return Path(result.value.path()).read_bytes()

            for width in (390,320):
                page.set_viewport_size({'width':width,'height':844})
                page.locator('#search').fill('Camera')
                page.locator('#select-visible-devices').check()
                for label in ('JSON','CSV','XLSX','PDF'):
                    data=download('#inventory',label,'selected')
                    if label=='JSON':assert {r['name'] for r in json.loads(data)['devices']}=={'Camera Alpha','Camera B'}
                    elif label=='CSV':assert {r['name'] for r in csv.DictReader(io.StringIO(data.decode('utf-8-sig')))}=={'Camera Alpha','Camera B'}
                    elif label=='XLSX':
                        book=load_workbook(io.BytesIO(data));assert book.sheetnames==['Devices'];assert book.active.max_row==3
                    else:
                        assert data.startswith(b'%PDF')
                        text=subprocess.check_output(['pdftotext','-','-'],input=data).decode()
                        assert 'Camera Alpha' in text and 'Camera B' in text and 'Audio desk' not in text
                assert page.evaluate('document.documentElement.scrollWidth<=innerWidth')
            page.locator('#delete-selected-devices').click()
            expect(page.locator('#batch-delete-title')).to_have_text('Delete 2 devices?')
            assert page.locator('#batch-delete-list li').count()==2
            page.locator('#cancel-batch-delete').click()
            assert len([r for r in app.inventory() if r['record_type']=='device'])==3
            page.locator('#delete-selected-devices').click();page.locator('#confirm-batch-delete').click()
            page.locator('#batch-delete-dialog').wait_for(state='hidden')
            assert [r['name'] for r in app.inventory() if r['record_type']=='device']==['Audio desk']
            open_workspace(page,'iptv-tab')
            expect(page.locator('.device-row')).to_have_count(2)
            page.locator('#source-buttons').get_by_role('button',name='Satellite',exact=True).click()
            assert page.locator('.device-row').count()==1
            assert page.locator('.device-row img').count()==0
            assert page.locator('.udp-vlc svg').count()==1
            assert [r['name'] for r in json.loads(download('#inventory','JSON','filtered'))['devices']]==['News stream']
            assert all(r['record_type']=='iptv' for r in json.loads(download('#inventory','JSON','all'))['devices'])
            open_workspace(page,'equipment-tab')
            page.wait_for_function("()=>!document.getElementById('equipment-inventory-select').disabled")
            page.locator('#equipment-inventory-select').select_option(str(inv['id']))
            page.wait_for_function("(id)=>document.getElementById('equipment-inventory-select').value===id&&!document.getElementById('equipment-inventory-select').disabled",arg=str(inv['id']))
            page.locator('#add-device').click()
            page.locator('#equipment-description').fill('HDMI converter')
            page.locator('#equipment-brand').fill('Generic')
            page.locator('#equipment-model').fill('HDMI-SDI')
            page.locator('#equipment-quantity').fill('1')
            page.locator('#equipment-notes').fill('With power supply')
            page.locator('#equipment-shelf').fill('B4')
            page.locator('#save-equipment').click()
            page.locator('#equipment-dialog').wait_for(state='hidden')
            expect(page.locator('#equipment-rows tr')).to_have_count(1)
            row=page.locator('#equipment-rows tr')
            row.get_by_role('button',name='New',exact=True).click()
            page.wait_for_function("()=>document.querySelector('#equipment-rows button[data-value=New]').getAttribute('aria-pressed')==='true'")
            row.locator('.equipment-edit-action').click()
            page.locator('#equipment-model').fill('Updated converter')
            page.locator('#save-equipment').click();page.locator('#equipment-dialog').wait_for(state='hidden')
            data=json.loads(download('#equipment-panel','JSON','all'))
            assert len(data['equipment'])==1
            saved=data['equipment'][0]
            assert saved['model']=='Updated converter'
            assert saved['custom_values']['unit_status']=='New'
            assert saved['notes']=='With power supply\nShelf: B4'
            row.locator('.equipment-delete-action').click()
            page.locator('#cancel-equipment-delete').click()
            assert len(app.equipment_inventory(inv['id']))==1
            row.locator('.equipment-delete-action').click();page.locator('#confirm-equipment-delete').click()
            page.locator('#equipment-delete-dialog').wait_for(state='hidden')
            assert app.equipment_inventory(inv['id'])==[]
            assert len(app.equipment_inventory(other['id']))==1
            assert page.locator('[aria-label="AVTrack version"]').inner_text()==version
            assert not errors,errors
            browser.close()
    finally:
        server.shutdown();server.server_close()
print('PASS: mobile keyboard/navigation, numeric commit-count version, CRUD, validation, confirmations, all report formats, tab/selection scope and safe deletion.')
