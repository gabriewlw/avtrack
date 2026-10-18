import base64
import copy
import io
import json
import zipfile
import unittest
from openpyxl import load_workbook
import app
import test_app
import workspace_transfer as transfer


class WorkspaceTransferTests(unittest.TestCase):
    setUpClass = test_app.AppTests.__dict__['setUpClass']
    tearDownClass = test_app.AppTests.__dict__['tearDownClass']
    setUp = test_app.AppTests.setUp
    tearDown = test_app.AppTests.tearDown
    request = test_app.AppTests.request

    def populate(self):
        app.save_device(dict(test_app.EXAMPLE))
        inventory=app.save_equipment_inventory({'name':'Spare Parts'})
        item=app.save_equipment(dict(inventory_id=inventory['id'],brand='Sony',model='X',description='Camera',serial_number='A',quantity=1,location='Locker',notes='Line one\nLine two'))
        app.set_equipment_review(item['id'],dict(item,data_checked=True))
        with app.connect() as con:
            con.execute("INSERT INTO import_preferences(target,settings) VALUES (?,?)",('device',json.dumps({'mappings':{'notes':{'name':'remarks'}}})))
        app.save_av_backup(dict(equipment_type='OBS profile',device_name='Control PC',filename='profile.json',notes='Recovery file',content=base64.b64encode(b'{"profile":"test"}').decode()))

    def test_complete_zip_and_json_round_trip_preserves_files_and_all_tables(self):
        self.populate()
        original=transfer.snapshot(app.connect,app.backup_directory)
        for kind in ('zip','json'):
            content,_,_=transfer.pack(original,kind)
            with app.connect() as con:
                con.execute('DELETE FROM devices')
                con.execute('DELETE FROM equipment')
            data=transfer.unpack(content)
            result=transfer.restore(data,app.connect,app.backup_directory)
            self.assertEqual(result['restored']['devices'],1)
            self.assertEqual(app.equipment_inventory(original['tables']['equipment'][0]['inventory_id'])[0]['notes'],'Line one\nLine two')
            restored=transfer.snapshot(app.connect,app.backup_directory)
            for name in original['tables']:
                if name!='av_backups':self.assertEqual(restored['tables'][name],original['tables'][name],name)
            self.assertEqual(sorted(restored['files'].values()),sorted(original['files'].values()))
            self.assertEqual(app.av_backup_file(1)[1],b'{"profile":"test"}')

    def test_invalid_backup_rolls_back_and_preserves_existing_files(self):
        self.populate()
        data=transfer.snapshot(app.connect,app.backup_directory)
        original_files=set(app.backup_directory().iterdir())
        data['tables']['devices'].append(dict(data['tables']['devices'][0]))
        with self.assertRaises(ValueError):
            transfer.restore(data,app.connect,app.backup_directory)
        self.assertEqual(len(app.inventory()),1)
        self.assertEqual(set(app.backup_directory().iterdir()),original_files)
        self.assertEqual(app.av_backup_file(1)[1],b'{"profile":"test"}')

    def test_reports_include_every_workspace_and_binary_files_are_not_cells(self):
        self.populate()
        data=transfer.snapshot(app.connect,app.backup_directory)
        content,_,_=transfer.pack(data,'xlsx')
        book=load_workbook(io.BytesIO(content))
        self.assertTrue({'Devices','IPTV','Spare Parts','Equipment backups'}.issubset(book.sheetnames))
        self.assertNotIn('import_preferences',book.sheetnames)
        self.assertNotIn('Content',[cell.value for cell in book['Equipment backups'][1]])
        content,_,_=transfer.pack(data,'csv')
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertIn('Devices.csv',archive.namelist())
            self.assertIn('Spare Parts.csv',archive.namelist())
        with self.assertRaises(ValueError):transfer.unpack(content)

    def test_http_export_restore_requires_confirmation_and_blocks_other_origins(self):
        from urllib.request import Request, urlopen
        from urllib.error import HTTPError
        self.populate()
        with urlopen(self.url+'/api/workspace/export?format=zip') as response:
            content=response.read()
            self.assertRegex(response.headers['Content-Disposition'],r'avtrack-all-data-.*\.zip')
        def restore(headers):
            request=Request(self.url+'/api/workspace/restore',data=content,method='POST',headers=headers)
            try:response=urlopen(request)
            except HTTPError as error:response=error
            with response:return response.status,json.load(response)
        self.assertEqual(restore({})[0],400)
        self.assertEqual(restore({'X-AVTrack-Restore':'replace-all','Origin':'https://other.example'})[0],403)
        status,result=restore({'X-AVTrack-Restore':'replace-all','Content-Type':'application/octet-stream'})
        self.assertEqual(status,200)
        self.assertEqual(result['restored']['devices'],1)

    def test_restore_can_be_retried_without_mutating_the_backup(self):
        self.populate()
        data=transfer.snapshot(app.connect,app.backup_directory)
        original=copy.deepcopy(data)
        transfer.restore(data,app.connect,app.backup_directory)
        self.assertEqual(data,original)
        transfer.restore(data,app.connect,app.backup_directory)
        self.assertEqual(app.av_backup_file(1)[1],b'{"profile":"test"}')

    def test_malformed_records_and_json_cannot_replace_working_data(self):
        self.populate()
        original=transfer.snapshot(app.connect,app.backup_directory)
        for table,field,value in [('equipment','quantity','many'),('equipment','quantity',-1),
                                  ('equipment','custom_values','[]'),('equipment_inventories','layout','{"columns":null}'),
                                  ('import_preferences','settings','not json')]:
            with self.subTest(table=table,field=field,value=value):
                bad=copy.deepcopy(original)
                bad['tables'][table][-1][field]=value
                with self.assertRaises(ValueError):transfer.restore(bad,app.connect,app.backup_directory)
                self.assertEqual(transfer.snapshot(app.connect,app.backup_directory)['tables'],original['tables'])
                self.assertEqual(transfer.snapshot(app.connect,app.backup_directory)['files'],original['files'])

    def test_truncated_attachment_cannot_replace_working_data(self):
        self.populate()
        data=transfer.snapshot(app.connect,app.backup_directory)
        data['files'][next(iter(data['files']))]=base64.b64encode(b'truncated').decode()
        with self.assertRaises(ValueError):transfer.restore(data,app.connect,app.backup_directory)
        self.assertEqual(app.av_backup_file(1)[1],b'{"profile":"test"}')

    def test_checksum_detects_changed_export_contents_and_legacy_backups_still_work(self):
        self.populate()
        original=transfer.snapshot(app.connect,app.backup_directory)
        content,_,_=transfer.pack(original,'json')
        changed=json.loads(content)
        changed['tables']['devices'][0]['name']='Unexpected change'
        with self.assertRaises(ValueError):transfer.unpack(json.dumps(changed).encode())
        # Existing backups from before integrity checks remain importable.
        transfer.restore(transfer.unpack(json.dumps(original).encode()),app.connect,app.backup_directory)
        self.assertEqual(app.inventory()[0]['name'],test_app.EXAMPLE['name'])

    def test_bad_archive_schema_and_paths_rejected(self):
        self.populate()
        data=transfer.snapshot(app.connect,app.backup_directory)
        bad=copy.deepcopy(data);bad['columns']={}
        with self.assertRaises(ValueError):transfer.restore(bad,app.connect,app.backup_directory)
        bad=copy.deepcopy(data);bad['files']['../escape']='eA=='
        with self.assertRaises(ValueError):transfer.restore(bad,app.connect,app.backup_directory)
        with self.assertRaises(ValueError):transfer.unpack(b'not a backup')
